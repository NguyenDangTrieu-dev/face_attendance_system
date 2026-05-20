from flask import Blueprint, jsonify, request
from psycopg2.extras import RealDictCursor

from db import get_db_connection
from auth import role_required

admin_bp = Blueprint('admin', __name__, url_prefix='/api/admin')


# ================= HELPERS =================
def table_exists(cur, table_name):
    cur.execute("""
        SELECT EXISTS (
            SELECT 1
            FROM information_schema.tables
            WHERE table_schema = 'public'
              AND table_name = %s
        ) AS exists
    """, (table_name,))
    row = cur.fetchone()
    return bool(row[0] if not isinstance(row, dict) else row["exists"])


def column_exists(cur, table_name, column_name):
    cur.execute("""
        SELECT EXISTS (
            SELECT 1
            FROM information_schema.columns
            WHERE table_schema = 'public'
              AND table_name = %s
              AND column_name = %s
        ) AS exists
    """, (table_name, column_name))
    row = cur.fetchone()
    return bool(row[0] if not isinstance(row, dict) else row["exists"])


def ensure_user_profile_columns(cur):
    """Các cột bổ sung dùng cho sinh viên. ADD COLUMN IF NOT EXISTS an toàn khi đã có cột."""
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS admission_course VARCHAR(20)")
    cur.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS major VARCHAR(120)")


def course_select_sql(cur):
    optional_cols = []
    for col in ("room", "period", "credits"):
        if column_exists(cur, "courses", col):
            optional_cols.append(f"c.{col}")
        else:
            optional_cols.append(f"NULL AS {col}")

    return f"""
        SELECT
            c.id,
            c.name,
            c.semester,
            c.lecturer_id,
            {optional_cols[0]},
            {optional_cols[1]},
            {optional_cols[2]},
            u.full_name AS lecturer_name
        FROM courses c
        LEFT JOIN users u
            ON c.lecturer_id = u.id
    """


def normalize_header(value):
    return str(value or "").strip().lower().replace(" ", "_")


def first_value(row, keys):
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip() != "":
            return str(value).strip()
    return ""


# ================= USERS =================
@admin_bp.route("/users", methods=["GET"])
@role_required([1])
def get_users():
    conn = None
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        ensure_user_profile_columns(cur)
        conn.commit()

        cur.execute("""
            SELECT id, full_name, role_id, admission_course, major
            FROM users
            ORDER BY id
        """)
        users = cur.fetchall()
        return jsonify(users)
    except Exception as e:
        return jsonify({"error": "Lỗi tải danh sách người dùng", "detail": str(e)}), 500
    finally:
        if conn:
            conn.close()


@admin_bp.route("/users", methods=["POST"])
@role_required([1])
def create_user():
    data = request.get_json(silent=True) or {}
    conn = None

    try:
        conn = get_db_connection()
        cur = conn.cursor()
        ensure_user_profile_columns(cur)

        role_id = int(data["role_id"])
        admission_course = data.get("admission_course") if role_id == 3 else None
        major = data.get("major") if role_id == 3 else None

        cur.execute("""
            INSERT INTO users (id, full_name, password, role_id, admission_course, major)
            VALUES (%s, %s, %s, %s, %s, %s)
        """, (
            data["id"],
            data["full_name"],
            data.get("password") or "123456",
            role_id,
            admission_course,
            major,
        ))

        conn.commit()
        return jsonify({"message": "User created"})

    except Exception as e:
        if conn:
            conn.rollback()
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ================= COURSES =================
@admin_bp.route("/courses", methods=["GET"])
@role_required([1])
def get_courses():
    conn = None
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)

        lecturer_id = request.args.get("lecturer_id")
        sql = course_select_sql(cur)
        params = []
        if lecturer_id:
            sql += " WHERE c.lecturer_id = %s"
            params.append(lecturer_id)
        sql += " ORDER BY c.id DESC"

        cur.execute(sql, params)
        data = cur.fetchall()
        return jsonify(data)

    except Exception as e:
        return jsonify({"error": "Lỗi tải danh sách học phần", "detail": str(e)}), 500
    finally:
        if conn:
            conn.close()


@admin_bp.route("/attendance/recent", methods=["GET"])
@role_required([1])
def get_recent_attendance():
    conn = None
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("""
            SELECT *
            FROM attendance
            ORDER BY time DESC
            LIMIT 100
        """)
        data = cur.fetchall()
        return jsonify(data)
    except Exception as e:
        return jsonify({"error": "Lỗi tải điểm danh gần đây", "detail": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ================= IMPORT LECTURERS =================
@admin_bp.route("/lecturers/import", methods=["POST"])
@role_required([1])
def import_lecturers():
    file = request.files.get("file")
    if not file:
        return jsonify({"error": "Chưa upload file Excel"}), 400

    try:
        from openpyxl import load_workbook
    except Exception:
        return jsonify({"error": "Thiếu thư viện openpyxl để đọc Excel"}), 500

    conn = None
    try:
        wb = load_workbook(file, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File Excel rỗng"}), 400

        headers = [normalize_header(h) for h in rows[0]]
        parsed_rows = []
        for raw in rows[1:]:
            item = {headers[i]: raw[i] for i in range(min(len(headers), len(raw)))}
            lecturer_id = first_value(item, ["lecturer_id", "teacher_id", "id", "ma_gv", "ma_giang_vien", "magv"])
            full_name = first_value(item, ["full_name", "name", "ho_ten", "hoten", "ten_giang_vien"])
            password = first_value(item, ["password", "mat_khau", "matkhau"])
            if lecturer_id and full_name:
                parsed_rows.append({
                    "id": lecturer_id,
                    "full_name": full_name,
                    "password": password or "123456",
                })

        if not parsed_rows:
            return jsonify({
                "error": "Không tìm thấy dòng hợp lệ. Cần cột lecturer_id/id và full_name/name."
            }), 400

        conn = get_db_connection()
        cur = conn.cursor()
        ensure_user_profile_columns(cur)

        created_or_updated = 0
        for item in parsed_rows:
            cur.execute("SELECT id FROM users WHERE id = %s", (item["id"],))
            existed = cur.fetchone() is not None
            if existed:
                cur.execute("""
                    UPDATE users
                    SET full_name = %s,
                        role_id = 2
                    WHERE id = %s
                """, (item["full_name"], item["id"]))
                if item["password"]:
                    cur.execute("UPDATE users SET password = %s WHERE id = %s", (item["password"], item["id"]))
            else:
                cur.execute("""
                    INSERT INTO users (id, full_name, password, role_id)
                    VALUES (%s, %s, %s, 2)
                """, (item["id"], item["full_name"], item["password"]))
            created_or_updated += 1

        conn.commit()
        return jsonify({
            "message": "Import giảng viên thành công",
            "count": created_or_updated,
        })

    except Exception as e:
        if conn:
            conn.rollback()
        return jsonify({"error": "Import giảng viên thất bại", "detail": str(e)}), 500
    finally:
        if conn:
            conn.close()


# ================= DELETE USER =================
@admin_bp.route("/users/<user_id>", methods=["DELETE"])
@role_required([1])
def delete_user(user_id):
    conn = None
    try:
        conn = get_db_connection()
        cur = conn.cursor()

        cur.execute("DELETE FROM course_students WHERE student_id = %s", (user_id,))

        if table_exists_for_delete(cur, "embeddings"):
            cur.execute("DELETE FROM embeddings WHERE user_id = %s", (user_id,))

        cur.execute("DELETE FROM users WHERE id = %s", (user_id,))

        conn.commit()
        return jsonify({"message": "Đã xóa người dùng"})

    except Exception as e:
        if conn:
            conn.rollback()
        print("DELETE USER ERROR:", e)
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            conn.close()


def table_exists_for_delete(cur, table_name):
    cur.execute("""
        SELECT EXISTS (
            SELECT 1
            FROM information_schema.tables
            WHERE table_schema = 'public'
              AND table_name = %s
        )
    """, (table_name,))
    return bool(cur.fetchone()[0])


# ================= UPDATE USER =================
@admin_bp.route("/users/<user_id>", methods=["PUT"])
@role_required([1])
def update_user(user_id):
    conn = None
    try:
        data = request.get_json(silent=True) or {}
        conn = get_db_connection()
        cur = conn.cursor()
        ensure_user_profile_columns(cur)

        role_id = int(data["role_id"])
        admission_course = data.get("admission_course") if role_id == 3 else None
        major = data.get("major") if role_id == 3 else None

        cur.execute("""
            UPDATE users
            SET full_name = %s,
                role_id = %s,
                admission_course = %s,
                major = %s
            WHERE id = %s
        """, (
            data["full_name"],
            role_id,
            admission_course,
            major,
            user_id,
        ))

        if data.get("password"):
            cur.execute("UPDATE users SET password = %s WHERE id = %s", (data["password"], user_id))

        if role_id != 3:
            cur.execute("DELETE FROM course_students WHERE student_id = %s", (user_id,))

        conn.commit()
        return jsonify({"message": "Đã cập nhật"})

    except Exception as e:
        if conn:
            conn.rollback()
        print("UPDATE USER ERROR:", e)
        return jsonify({"error": str(e)}), 500
    finally:
        if conn:
            conn.close()
